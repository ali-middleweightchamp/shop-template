"""
Разбор и применение прайса. Устойчив к «кривым» файлам: гибкие заголовки
(ru/uz, любой регистр), цены вида «12 500 сум», наличие да/нет/+/-/bor/yo'q.
Битые строки не роняют импорт — попадают в отчёт с номером строки.
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.text import slugify

from apps.catalog.models import Category, Product, unique_slug

# --- Сопоставление заголовков (нормализованных) с полями ---
HEADER_ALIASES = {
    "sku": {"артикул", "sku", "artikul", "код", "kod"},
    "name": {"название", "наименование", "name", "nomi", "nomlanishi", "tovar"},
    "name_uz": {"названиеuz", "nomiuz", "nameuz", "nomuz", "названиеузб", "узб"},
    "category": {"категория", "category", "turkum", "turkumi", "kategoriya", "раздел"},
    "price": {"цена", "price", "narx", "narxi", "стоимость"},
    "unit": {"единица", "ед", "едизм", "unit", "birlik", "o'lchov"},
    "in_stock": {"вналичии", "наличие", "stock", "instock", "mavjud", "mavjudlik", "bor"},
    "description": {"описание", "description", "tavsif", "izoh"},
}

UNIT_ALIASES = {
    "шт": "шт", "штука": "шт", "pcs": "шт", "dona": "шт", "piece": "шт",
    "упак": "упак", "уп": "упак", "pack": "упак", "o'ram": "упак", "quti": "упак", "paket": "упак",
    "кг": "кг", "kg": "кг",
    "м": "м", "m": "м", "метр": "м",
}

STOCK_TRUE = {"да", "yes", "+", "1", "true", "bor", "есть", "ha", "в наличии", "вналичии"}
STOCK_FALSE = {"нет", "no", "-", "0", "false", "yo'q", "yoq", "yoq'", "нет в наличии"}

CANONICAL_HEADERS = ["Артикул", "Название", "Название (uz)", "Категория", "Цена", "Единица", "В наличии", "Описание"]
DEFAULT_CATEGORY = "Без категории"


def _norm(s):
    """Нормализация ключа: нижний регистр, без пробелов/точек/скобок."""
    return re.sub(r"[\s.()\-_/]", "", str(s or "").strip().lower())


def _header_map(headers):
    """{индекс колонки: имя поля} по нормализованным заголовкам."""
    mapping = {}
    for idx, h in enumerate(headers):
        key = _norm(h)
        for field, aliases in HEADER_ALIASES.items():
            if key in aliases and field not in mapping.values():
                mapping[idx] = field
                break
    return mapping


def parse_price(value):
    """«12 500 сум», «12500,00», «12 500» → Decimal. Иначе ValueError."""
    s = str(value if value is not None else "").replace("\xa0", " ")
    s = re.sub(r"[^\d,.\s-]", "", s).replace(" ", "")
    if not s:
        raise ValueError("пусто")
    if "," in s and "." in s:
        s = s.replace(",", "")  # запятая — разделитель тысяч
    elif "," in s:
        s = s.replace(",", ".")  # запятая — десятичный разделитель
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError("не число")
    if d < 0:
        raise ValueError("отрицательная")
    return d


def parse_stock(value, default=True):
    """да/нет, +/-, 1/0, true/false, bor/yo'q → bool."""
    s = str(value if value is not None else "").strip().lower().replace("’", "'").replace("‘", "'")
    if s == "":
        return default
    if s in STOCK_TRUE:
        return True
    if s in STOCK_FALSE:
        return False
    return default


def parse_unit(value):
    key = _norm(value)
    return UNIT_ALIASES.get(key, "шт")


def _read_rows(django_file):
    """Читает xlsx/xls/csv → (headers, list_of_row_lists)."""
    name = (getattr(django_file, "name", "") or "").lower()
    data = django_file.read()
    if name.endswith(".csv"):
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = data.decode("utf-8", errors="replace")
        # автоопределение разделителя , или ;
        sample = text[:2000]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [r for r in reader]
    else:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(str(c).strip() for c in r)]  # убрать пустые строки
    if not rows:
        return [], []
    return rows[0], rows[1:]


def analyze(django_file):
    """Разбор без записи в БД: что создастся/обновится/ошибки (с № строк)."""
    headers, body = _read_rows(django_file)
    field_by_col = _header_map(headers)
    fields = set(field_by_col.values())

    result = {
        "headers": headers,
        "rows_total": len(body),
        "to_create": 0,
        "to_update": 0,
        "errors": [],
        "valid": [],  # нормализованные строки для применения
        "missing_columns": [],
    }
    if "name" not in fields:
        result["missing_columns"].append("Название")
    if "price" not in fields:
        result["missing_columns"].append("Цена")
    if result["missing_columns"]:
        return result

    existing_skus = set(Product.objects.values_list("sku", flat=True))
    existing_names = dict(Product.objects.values_list("name_ru", "sku"))

    for i, raw in enumerate(body):
        row_no = i + 2  # +1 заголовок, +1 к человекочитаемому номеру
        rec = {}
        for col, field in field_by_col.items():
            rec[field] = raw[col] if col < len(raw) else ""

        name = str(rec.get("name", "")).strip()
        if not name:
            result["errors"].append({"row": row_no, "message": "нет названия"})
            continue
        try:
            price = parse_price(rec.get("price"))
        except ValueError:
            result["errors"].append(
                {"row": row_no, "message": f"цена не является числом: «{rec.get('price')}»"}
            )
            continue

        sku = str(rec.get("sku", "")).strip()
        item = {
            "sku": sku,
            "name": name,
            "name_uz": str(rec.get("name_uz", "")).strip(),
            "category": str(rec.get("category", "")).strip() or DEFAULT_CATEGORY,
            "price": price,
            "unit": parse_unit(rec.get("unit")),
            "in_stock": parse_stock(rec.get("in_stock")),
            "description": str(rec.get("description", "")).strip(),
            "row": row_no,
        }
        # создастся или обновится?
        if sku:
            (result.__setitem__("to_update", result["to_update"] + 1) if sku in existing_skus
             else result.__setitem__("to_create", result["to_create"] + 1))
        elif name in existing_names:
            result["to_update"] += 1
        else:
            result["to_create"] += 1
        result["valid"].append(item)

    return result


@transaction.atomic
def apply(django_file):
    """Применяет прайс в одной транзакции. Возвращает статистику."""
    data = analyze(django_file)
    if data["missing_columns"]:
        return data

    created = updated = 0
    cat_cache = {}

    for item in data["valid"]:
        cat_name = item["category"]
        category = cat_cache.get(cat_name.lower())
        if category is None:
            category, _ = Category.objects.get_or_create(
                name_ru=cat_name, defaults={"name": cat_name}
            )
            cat_cache[cat_name.lower()] = category

        sku = item["sku"]
        product = None
        if sku:
            product = Product.objects.filter(sku=sku).first()
        if product is None and not sku:
            product = Product.objects.filter(name_ru=item["name"]).first()

        if product is None:
            product = Product(sku=sku or _gen_sku(item["name"]))
            created += 1
        else:
            updated += 1

        product.name = item["name"]
        product.name_ru = item["name"]
        if item["name_uz"]:
            product.name_uz = item["name_uz"]
        product.category = category
        product.price = item["price"]
        product.unit = item["unit"]
        product.in_stock = item["in_stock"]
        if item["description"]:
            product.description = item["description"]
        product.is_active = True
        product.save()

    data["rows_created"] = created
    data["rows_updated"] = updated
    return data


def _gen_sku(name):
    base = (slugify(name)[:40] or "sku").upper().replace("-", "")
    return unique_slug(Product, base or "SKU").upper()


# --- Шаблон и экспорт ---
def build_template():
    """xlsx-шаблон с заголовками и примерами строк. Возвращает bytes."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(CANONICAL_HEADERS)
    ws.append(["BUM-A4", "Бумага А4, 500 л", "A4 qog'oz, 500 varaq", "Бумага", "48 500", "шт", "да", "Класс C"])
    ws.append(["RCH-01", "Ручка синяя", "Ko'k ruchka", "Ручки", "3200", "упак", "bor", ""])
    for col in range(1, len(CANONICAL_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_catalog():
    """Текущий каталог в xlsx. Возвращает bytes."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Каталог"
    ws.append(CANONICAL_HEADERS)
    for p in Product.objects.select_related("category").order_by("category__name", "name"):
        ws.append([
            p.sku, p.name_ru, p.name_uz or "", p.category.name_ru,
            int(p.price), p.unit, "да" if p.in_stock else "нет", p.description_ru or "",
        ])
    for col in range(1, len(CANONICAL_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
